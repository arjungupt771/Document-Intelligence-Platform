import pytest
from docx import Document as DocxDocument

from app.parsing.docx_parser import DOCXParser
from app.parsing.errors import CorruptedDocumentError, EmptyDocumentError


def _make_docx(tmp_path, paragraphs=None, heading=None, table_rows=None, filename="sample.docx"):
    document = DocxDocument()

    if heading:
        document.add_heading(heading, level=1)

    for paragraph_text in paragraphs or []:
        document.add_paragraph(paragraph_text)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value

    path = tmp_path / filename
    document.save(path)
    return path


def test_parses_paragraphs_and_heading(tmp_path):
    path = _make_docx(tmp_path, paragraphs=["Body text here."], heading="Purchase Order")

    parsed = DOCXParser().parse(path, document_id="doc-1", source_filename="sample.docx")

    assert parsed.parser_name == "python-docx"
    assert "Purchase Order" in parsed.text
    assert "Body text here." in parsed.text

    heading_blocks = [b for b in parsed.pages[0].text_blocks if b.block_type.value == "heading"]
    assert len(heading_blocks) == 1


def test_parses_tables(tmp_path):
    path = _make_docx(
        tmp_path,
        paragraphs=["intro"],
        table_rows=[["Item", "Qty"], ["Widget", "3"]],
    )

    parsed = DOCXParser().parse(path, document_id="doc-1", source_filename="sample.docx")

    assert len(parsed.tables) == 1
    table = parsed.tables[0]
    assert table.rows == 2
    assert table.columns == 2
    assert {cell.text for cell in table.cells} == {"Item", "Qty", "Widget", "3"}


def test_provenance_has_no_page_number_for_docx(tmp_path):
    path = _make_docx(tmp_path, paragraphs=["Only paragraph"])

    parsed = DOCXParser().parse(path, document_id="doc-1", source_filename="sample.docx")

    block = parsed.pages[0].text_blocks[0]
    assert block.provenance.page_number is None
    assert block.provenance.source_parser == "python-docx"


def test_corrupted_docx_raises_corrupted_document_error(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"this is not a real docx file")

    with pytest.raises(CorruptedDocumentError):
        DOCXParser().parse(path, document_id="doc-1", source_filename="bad.docx")


def test_empty_docx_raises_empty_document_error(tmp_path):
    path = _make_docx(tmp_path, filename="empty.docx")

    with pytest.raises(EmptyDocumentError):
        DOCXParser().parse(path, document_id="doc-1", source_filename="empty.docx")
