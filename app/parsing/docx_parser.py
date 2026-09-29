from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.parsing.base import Parser
from app.parsing.errors import CorruptedDocumentError, EmptyDocumentError
from app.parsing.models import (
    BlockType,
    DocumentMetadata,
    Page,
    ParsedDocument,
    Provenance,
    TableBlock,
    TableCell,
    TextBlock,
)

PARSER_NAME = "python-docx"

# .docx has no native page concept until it's laid out/rendered, so the
# whole document is modeled as a single logical page.
_SINGLE_PAGE_NUMBER = 1


class DOCXParser(Parser):
    """Adapter from python-docx documents to ParsedDocument."""

    supported_extensions = frozenset({".docx"})

    def parse(
        self,
        path: Path,
        *,
        document_id: str,
        source_filename: str,
    ) -> ParsedDocument:
        try:
            document = DocxDocument(str(path))
        except PackageNotFoundError as exc:
            raise CorruptedDocumentError(source_filename, "not a valid .docx package") from exc
        except Exception as exc:
            raise CorruptedDocumentError(source_filename, str(exc)) from exc

        text_blocks = self._extract_text_blocks(document)
        tables = self._extract_tables(document, start_index=len(text_blocks))

        page = Page(
            page_number=_SINGLE_PAGE_NUMBER,
            width=None,
            height=None,
            text_blocks=text_blocks,
            tables=tables,
            images=[],
        )

        parsed_document = ParsedDocument(
            document_id=document_id,
            source_filename=source_filename,
            parser_name=PARSER_NAME,
            metadata=self._build_metadata(document),
            pages=[page],
        )

        if not parsed_document.text.strip() and not parsed_document.tables:
            raise EmptyDocumentError(source_filename)

        return parsed_document

    def _extract_text_blocks(self, document: "DocxDocument") -> list[TextBlock]:
        text_blocks: list[TextBlock] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = (paragraph.style.name if paragraph.style else "") or ""
            block_type = (
                BlockType.HEADING if style_name.lower().startswith("heading") else BlockType.TEXT
            )

            text_blocks.append(
                TextBlock(
                    text=text,
                    block_type=block_type,
                    provenance=Provenance(
                        source_parser=PARSER_NAME,
                        page_number=None,
                        block_index=len(text_blocks),
                    ),
                )
            )

        return text_blocks

    def _extract_tables(self, document: "DocxDocument", start_index: int) -> list[TableBlock]:
        tables: list[TableBlock] = []

        for offset, table in enumerate(document.tables):
            cells = [
                TableCell(text=cell.text.strip(), row=row_index, column=column_index)
                for row_index, row in enumerate(table.rows)
                for column_index, cell in enumerate(row.cells)
            ]

            tables.append(
                TableBlock(
                    rows=len(table.rows),
                    columns=len(table.columns),
                    cells=cells,
                    provenance=Provenance(
                        source_parser=PARSER_NAME,
                        page_number=None,
                        block_index=start_index + offset,
                    ),
                )
            )

        return tables

    @staticmethod
    def _build_metadata(document: "DocxDocument") -> DocumentMetadata:
        core_properties = document.core_properties
        return DocumentMetadata(
            title=core_properties.title or None,
            author=core_properties.author or None,
            created_at=core_properties.created,
            modified_at=core_properties.modified,
            page_count=_SINGLE_PAGE_NUMBER,
            source_content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )
